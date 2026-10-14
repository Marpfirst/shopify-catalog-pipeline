"""Exports: Shopify product-import CSV, Excel report, local image assets."""

from __future__ import annotations

import csv
import sqlite3
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

SHOPIFY_COLUMNS = [
    "Handle", "Title", "Body (HTML)", "Vendor", "Type", "Tags", "Published",
    "Option1 Name", "Option1 Value", "Option2 Name", "Option2 Value", "Option3 Name", "Option3 Value",
    "Variant SKU", "Variant Grams", "Variant Inventory Policy", "Variant Fulfillment Service",
    "Variant Price", "Variant Compare At Price", "Variant Requires Shipping",
    "Image Src", "Image Position", "Variant Image", "Status",
]


def shopify_rows(con: sqlite3.Connection, supplier: str):
    """Shopify CSV layout: first row carries product fields, later rows only variants/images."""
    for p in con.execute("SELECT * FROM products WHERE supplier=? ORDER BY title", (supplier,)):
        names = (p["option_names"].split(" | ") + ["", "", ""])[:3]
        variants = con.execute("SELECT * FROM variants WHERE supplier=? AND product_id=?",
                               (supplier, p["product_id"])).fetchall()
        images = con.execute("SELECT * FROM images WHERE supplier=? AND product_id=? ORDER BY position",
                             (supplier, p["product_id"])).fetchall()
        for i in range(max(len(variants), len(images))):
            row = dict.fromkeys(SHOPIFY_COLUMNS, "")
            row["Handle"] = p["handle"]
            if i == 0:
                row.update({"Title": p["title"], "Body (HTML)": p["body_html"], "Vendor": p["vendor"],
                            "Type": p["product_type"], "Tags": p["tags"], "Published": "TRUE",
                            "Status": "draft"})
            if i < len(variants):
                v = variants[i]
                for n, name in enumerate(names, 1):
                    if name and v[f"option{n}"]:
                        if i == 0:
                            row[f"Option{n} Name"] = name
                        row[f"Option{n} Value"] = v[f"option{n}"]
                row.update({"Variant SKU": v["sku"], "Variant Grams": v["grams"],
                            "Variant Inventory Policy": "deny", "Variant Fulfillment Service": "manual",
                            "Variant Price": f"{v['price']:.2f}",
                            "Variant Compare At Price": f"{v['compare_at_price']:.2f}" if v["compare_at_price"] else "",
                            "Variant Requires Shipping": "TRUE", "Variant Image": v["image_src"]})
            if i < len(images):
                row.update({"Image Src": images[i]["src"], "Image Position": images[i]["position"]})
            yield row


def write_shopify_csv(con: sqlite3.Connection, supplier: str, path: Path) -> int:
    rows = list(shopify_rows(con, supplier))
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, SHOPIFY_COLUMNS)
        w.writeheader()
        w.writerows(rows)
    return len(rows)


def write_report(con: sqlite3.Connection, path: Path, mapping: list[dict] | None = None) -> None:
    wb = Workbook()
    sheets = {
        "Variants": """SELECT v.supplier, p.title AS product, v.title AS variant, v.sku, v.price,
                         v.compare_at_price, v.available, p.handle, v.variant_id, v.image_src
                       FROM variants v JOIN products p USING (supplier, product_id)
                       ORDER BY v.supplier, p.title""",
        "Changes": "SELECT * FROM changes ORDER BY at DESC",
    }
    wb.remove(wb.active)
    for name, sql in sheets.items():
        cur = con.execute(sql)
        _sheet(wb, name, [d[0] for d in cur.description], [list(r) for r in cur])
    if mapping:
        cols = list(dict.fromkeys(k for r in mapping for k in r))
        _sheet(wb, "Variant mapping", cols, [[r.get(c, "") for c in cols] for r in mapping])
    wb.save(path)


def _sheet(wb: Workbook, name: str, header: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(h), 12) + 2, 45)


def download_images(con: sqlite3.Connection, supplier: str, out_dir: Path, per_product: int = 3) -> int:
    """Save original-resolution images as <out>/<supplier>/<handle>/<position>.<ext>."""
    count = 0
    s = requests.Session()
    for p in con.execute("SELECT product_id, handle FROM products WHERE supplier=?", (supplier,)):
        for img in con.execute("SELECT position, src FROM images WHERE supplier=? AND product_id=? "
                               "ORDER BY position LIMIT ?", (supplier, p["product_id"], per_product)):
            ext = img["src"].split("?")[0].rsplit(".", 1)[-1][:4]
            dest = out_dir / supplier / p["handle"] / f"{img['position']}.{ext}"
            if dest.exists():
                continue
            dest.parent.mkdir(parents=True, exist_ok=True)
            resp = s.get(img["src"], timeout=60)
            resp.raise_for_status()
            dest.write_bytes(resp.content)
            count += 1
    return count
